"""Normalize Kentucky's official 2024-25 district ACT export.

The Kentucky workbook contains historical school rows and a district-total row.
This extracts only the current-year district totals and keeps the official
student count.  Unmatched entities are reported rather than guessed.
"""

import csv
import io
import zipfile
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "work/raw/ky/ACT_Average_20242025.xlsx"
NCES_ZIP = ROOT / "work/raw/ccd_lea_2023_24.zip"
OUTPUT = ROOT / "work/raw/ky/ky_act_performance_observations.csv"


def clean(value):
    if value is None:
        return ""
    value = str(value).strip().rstrip(".")
    return "" if value in {"", "*", "N/A", "."} else value


def main() -> None:
    with zipfile.ZipFile(NCES_ZIP) as archive:
        nces_csv = next(name for name in archive.namelist() if name.lower().endswith(".csv"))
        rows = csv.DictReader(io.TextIOWrapper(archive.open(nces_csv), encoding="utf-8-sig"))
        nces_by_name = {
            row["LEA_NAME"].strip().upper(): row["LEAID"]
            for row in rows
            if row["ST"] == "KY"
        }

    workbook = load_workbook(INPUT, read_only=True, data_only=True)
    sheet = workbook["AVG_SCORE"]
    observations = []
    unmatched = set()
    headers = [cell.value for cell in next(sheet.iter_rows())]
    for values in sheet.iter_rows(min_row=2, values_only=True):
        row = dict(zip(headers, values))
        if row["Year"] != "2024-25" or row["School Name"] != "-- District Total --":
            continue
        district_name = clean(row["District Name"])
        if district_name == "STATE":
            continue
        district_id = nces_by_name.get(district_name.upper())
        if not district_id:
            unmatched.add(district_name)
            continue
        score = clean(row["Composite Average Score"])
        if not score:
            continue
        observations.append({
            "observation_id": f"ky_{district_id}_act_2024_25",
            "nces_school_id": "",
            "nces_district_id": district_id,
            "test_type": "ACT",
            "metric_name": "average_composite",
            "metric_value": score,
            "score_year": "2024-25",
            "students_tested": clean(row["Number of Students"]),
            "participation_rate": "",
            "aggregation_method": "official_district_total_grade_11_average",
            "source_type": "official_state_export",
            "source_id": "kentucky_act_2024_25_scores",
            "confidence": "high",
            "notes": "Kentucky district total; ACT grade 11 average score by site; all students aggregate as published.",
        })

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    columns = list(observations[0].keys())
    with OUTPUT.open("w", newline="") as destination:
        writer = csv.DictWriter(destination, fieldnames=columns)
        writer.writeheader()
        writer.writerows(sorted(observations, key=lambda item: item["nces_district_id"]))
    print(f"Wrote {len(observations):,} Kentucky district ACT observations to {OUTPUT}")
    print(f"Excluded {len(unmatched):,} unmatched Kentucky entities: {sorted(unmatched)}")


if __name__ == "__main__":
    main()
